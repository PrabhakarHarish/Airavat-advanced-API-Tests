import io
import zipfile
import openpyxl
from openpyxl.styles import Font


def create_valid_financial_workbook() -> io.BytesIO:
    wb = openpyxl.Workbook()

    ws_is = wb.active
    ws_is.title = "Income Statement"
    ws_is["A1"] = "Airavat Base Financial Template"
    ws_is["A1"].font = Font(bold=True, size=14)

    ws_is["A3"] = "Period"
    ws_is["B3"] = "FY2024"
    ws_is["C3"] = "FY2025"
    ws_is["D3"] = "FY2026"

    ws_is["A16"] = "Total Operating Revenue"
    ws_is["D16"] = 15000000

    ws_is["A17"] = "Total Operating Expenses"
    ws_is["D17"] = 10000000

    ws_is["A18"] = "Operating Income (EBITDA)"
    ws_is["D18"] = "=D16-D17"

    ws_bs = wb.create_sheet(title="Balance Sheet")
    ws_bs["A1"] = "Balance Sheet"
    ws_bs["A3"] = "Total Assets"
    ws_bs["B3"] = 50000000
    ws_bs["A4"] = "Total Liabilities & Equity"
    ws_bs["B4"] = 50000000

    ws_cf = wb.create_sheet(title="Cash Flow")
    ws_cf["A1"] = "Cash Flow Statement"
    ws_cf["A3"] = "Operating Cash Flow"
    ws_cf["B3"] = 4500000

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def create_corrupted_financial_workbook() -> io.BytesIO:
    valid_buf = create_valid_financial_workbook()
    in_zip = zipfile.ZipFile(valid_buf, 'r')

    out_buf = io.BytesIO()
    with zipfile.ZipFile(out_buf, 'w', zipfile.ZIP_DEFLATED) as out_zip:
        for item in in_zip.infolist():
            data = in_zip.read(item.filename)
            if "sheet2.xml" in item.filename:
                data = b"<worksheet><sheetData><row r='42'><c r='A42'><corrupt><<<<BROKEN_XML_TAGS"
            out_zip.writestr(item, data)

    out_buf.seek(0)
    return out_buf


def create_zip_bomb_xlsx() -> io.BytesIO:
    out_buf = io.BytesIO()
    with zipfile.ZipFile(out_buf, 'w', zipfile.ZIP_DEFLATED) as z:
        huge_zero_data = b"0" * (10 * 1024 * 1024)
        z.writestr("xl/worksheets/sheet1.xml", huge_zero_data)
        z.writestr("[Content_Types].xml", b"<?xml version='1.0' encoding='UTF-8'?><Types></Types>")
    out_buf.seek(0)
    return out_buf


def create_xxe_xlsx() -> io.BytesIO:
    xxe_workbook_xml = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        b'<!DOCTYPE foo [ <!ENTITY xxe SYSTEM "file:///etc/passwd"> ]>\n'
        b'<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        b'<sheets><sheet name="&xxe;" sheetId="1" r:id="rId1"/></sheets>'
        b'</workbook>'
    )
    valid_buf = create_valid_financial_workbook()
    in_zip = zipfile.ZipFile(valid_buf, 'r')

    out_buf = io.BytesIO()
    with zipfile.ZipFile(out_buf, 'w', zipfile.ZIP_DEFLATED) as out_zip:
        for item in in_zip.infolist():
            if item.filename == "xl/workbook.xml":
                out_zip.writestr(item, xxe_workbook_xml)
            else:
                out_zip.writestr(item, in_zip.read(item.filename))

    out_buf.seek(0)
    return out_buf


def create_macro_enabled_xlsx() -> io.BytesIO:
    valid_buf = create_valid_financial_workbook()
    in_zip = zipfile.ZipFile(valid_buf, 'r')

    out_buf = io.BytesIO()
    with zipfile.ZipFile(out_buf, 'w', zipfile.ZIP_DEFLATED) as out_zip:
        for item in in_zip.infolist():
            out_zip.writestr(item, in_zip.read(item.filename))
        out_zip.writestr("xl/vbaProject.bin", b"MZ_SIMULATED_MALICIOUS_VBA_PAYLOAD_MACRO")

    out_buf.seek(0)
    return out_buf


def create_formula_injection_xlsx() -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "=cmd|' /C calc'!A0"
    ws["A2"] = "@SUM(1+1)*cmd|' /C calc'!A0"
    ws["B1"] = 1000
    ws["B2"] = 2000

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def create_truncated_header_file() -> bytes:
    return b"PK\x03\x04\x14\x00\x00\x00\x08\x00"
